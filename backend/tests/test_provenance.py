"""Golden-file structural validity tests for all six Provenance Catalog profiles."""

from __future__ import annotations

import csv
import email.parser
import io

import pytest

from evidence_factory.provenance.catalog import ProvenanceCatalog


@pytest.fixture()
def catalog() -> ProvenanceCatalog:
    return ProvenanceCatalog()


# ---------------------------------------------------------------------------
# Email (.eml) — RFC822 round-trip
# ---------------------------------------------------------------------------


def test_email_is_valid_rfc822(catalog: ProvenanceCatalog) -> None:
    content = "Hi Alice, let's schedule the budget review for next Tuesday."
    metadata = {
        "from": "bob@example.com",
        "to": "alice@example.com",
        "subject": "Budget Review",
        "message_id": "<test001@evidence-factory.local>",
        "date": "Mon, 15 Jan 2024 09:30:00 +0000",
    }
    from evidence_factory.models import ArtifactProfile as _AP

    raw = catalog.write(
        profile=_AP.EMAIL,
        text_content=content,
        metadata=metadata,
    )
    assert isinstance(raw, bytes)
    assert len(raw) > 0

    # Parse with stdlib
    parser = email.parser.BytesParser()
    msg = parser.parsebytes(raw)
    assert msg["From"] == "bob@example.com"
    assert msg["To"] == "alice@example.com"
    assert msg["Subject"] == "Budget Review"
    body = msg.get_payload()
    assert content[:20] in (body or "")


def test_email_sha256_stable_on_identical_payload(catalog: ProvenanceCatalog) -> None:
    import hashlib

    from evidence_factory.models import ArtifactProfile

    metadata = {
        "from": "alice@example.com",
        "to": "bob@example.com",
        "subject": "Stable test",
        "message_id": "<stable@evidence-factory.local>",
        "date": "Mon, 15 Jan 2024 12:00:00 +0000",
    }
    raw1 = catalog.write(ArtifactProfile.EMAIL, "Same body text.", metadata)
    raw2 = catalog.write(ArtifactProfile.EMAIL, "Same body text.", metadata)
    assert hashlib.sha256(raw1).hexdigest() == hashlib.sha256(raw2).hexdigest()


# ---------------------------------------------------------------------------
# SMS / chat (CSV)
# ---------------------------------------------------------------------------


def test_sms_is_valid_csv(catalog: ProvenanceCatalog) -> None:
    from evidence_factory.models import ArtifactProfile

    content = "Are you free for lunch?\nYes, see you at noon."
    metadata = {
        "sender": "alice",
        "recipient": "bob",
        "timestamp": "2024-01-15 12:00:00",
    }
    raw = catalog.write(ArtifactProfile.SMS, content, metadata)
    assert isinstance(raw, bytes)
    decoded = raw.decode("utf-8")
    reader = csv.DictReader(io.StringIO(decoded))
    rows = list(reader)
    assert len(rows) >= 1
    assert "sender" in reader.fieldnames  # type: ignore[operator]
    assert "message" in reader.fieldnames  # type: ignore[operator]


# ---------------------------------------------------------------------------
# PDF document
# ---------------------------------------------------------------------------


def test_pdf_starts_with_header(catalog: ProvenanceCatalog) -> None:
    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(
        ArtifactProfile.PDF,
        "This is a test document.",
        {"author": "Alice Smith", "created": "D:20240115090000"},
    )
    assert isinstance(raw, bytes)
    assert raw.startswith(b"%PDF-"), f"Expected PDF header, got {raw[:8]!r}"


def test_pdf_contains_eof_marker(catalog: ProvenanceCatalog) -> None:
    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(ArtifactProfile.PDF, "Document body.", {"author": "Bob"})
    assert b"%%EOF" in raw


def test_pdf_contains_text_content(catalog: ProvenanceCatalog) -> None:
    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(ArtifactProfile.PDF, "FindThisInPDF", {"author": "Test"})
    # The text is embedded in the content stream (latin-1 encoded)
    assert b"FindThisInPDF" in raw


# ---------------------------------------------------------------------------
# XLSX ledger
# ---------------------------------------------------------------------------


def test_xlsx_is_valid_excel(catalog: ProvenanceCatalog) -> None:
    pytest.importorskip("openpyxl")
    import openpyxl

    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(
        ArtifactProfile.XLSX,
        "Revenue Q1: $1,200\nRevenue Q2: $1,800",
        {"author": "Alice", "sheet_title": "Ledger", "timestamp": "2024-01-15"},
    )
    assert isinstance(raw, bytes)
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    assert len(wb.sheetnames) >= 1


def test_xlsx_has_data_rows(catalog: ProvenanceCatalog) -> None:
    pytest.importorskip("openpyxl")
    import openpyxl

    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(
        ArtifactProfile.XLSX,
        "Line one\nLine two\nLine three",
        {"author": "Bob", "timestamp": "2024-01-16"},
    )
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb.active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    # Header row + data rows
    assert len(rows) >= 2


# ---------------------------------------------------------------------------
# JPEG photo (Pillow + piexif)
# ---------------------------------------------------------------------------


def test_jpeg_starts_with_magic_bytes(catalog: ProvenanceCatalog) -> None:
    pytest.importorskip("PIL")
    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(
        ArtifactProfile.JPEG,
        "Photo caption here.",
        {"datetime": "2024:01:15 10:30:00", "device": "alice_phone"},
    )
    assert isinstance(raw, bytes)
    assert raw[:2] == b"\xff\xd8", f"Expected JPEG magic, got {raw[:2]!r}"


def test_jpeg_contains_exif(catalog: ProvenanceCatalog) -> None:
    pytest.importorskip("piexif")
    pytest.importorskip("PIL")
    import piexif

    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(
        ArtifactProfile.JPEG,
        "A photo of the office.",
        {"datetime": "2024:01:15 10:30:00", "device": "cam01"},
    )
    exif = piexif.load(raw)
    # DateTimeOriginal should be set
    dt = exif["Exif"].get(piexif.ExifIFD.DateTimeOriginal)
    assert dt is not None


def test_jpeg_with_gps(catalog: ProvenanceCatalog) -> None:
    pytest.importorskip("piexif")
    pytest.importorskip("PIL")
    import piexif

    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(
        ArtifactProfile.JPEG,
        "Park photo.",
        {"datetime": "2024:01:20 14:00:00", "device": "cam01", "gps_lat": 51.5, "gps_lon": -0.1},
    )
    exif = piexif.load(raw)
    assert piexif.GPSIFD.GPSLatitude in exif["GPS"]


def test_jpeg_make_model_from_metadata(catalog: ProvenanceCatalog) -> None:
    """Make and Model EXIF tags must come from metadata, not be hardcoded."""
    pytest.importorskip("piexif")
    pytest.importorskip("PIL")
    import piexif

    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(
        ArtifactProfile.JPEG,
        "Device photo.",
        {"datetime": "2024:01:15 10:30:00", "make": "Apple", "model": "iPhone 15 Pro"},
    )
    exif = piexif.load(raw)
    make = exif["0th"].get(piexif.ImageIFD.Make, b"").decode()
    model = exif["0th"].get(piexif.ImageIFD.Model, b"").decode()
    assert "Apple" in make
    assert "iPhone 15 Pro" in model


def test_jpeg_datetime_original_matches_metadata(catalog: ProvenanceCatalog) -> None:
    """DateTimeOriginal EXIF tag must match the datetime provided in metadata."""
    pytest.importorskip("piexif")
    pytest.importorskip("PIL")
    import piexif

    from evidence_factory.models import ArtifactProfile

    dt_str = "2024:03:15 14:22:00"
    raw = catalog.write(
        ArtifactProfile.JPEG,
        "Timestamp test.",
        {"datetime": dt_str, "device": "cam01"},
    )
    exif = piexif.load(raw)
    dto = exif["Exif"].get(piexif.ExifIFD.DateTimeOriginal, b"").decode()
    assert dto == dt_str


def test_jpeg_exif_round_trip_two_independent_libraries(catalog: ProvenanceCatalog) -> None:
    """EXIF data read by piexif and exifread must agree on DateTimeOriginal."""
    pytest.importorskip("piexif")
    pytest.importorskip("PIL")
    pytest.importorskip("exifread")
    import io as _io

    import exifread
    import piexif

    from evidence_factory.models import ArtifactProfile

    dt_str = "2024:06:01 09:00:00"
    raw = catalog.write(
        ArtifactProfile.JPEG,
        "Two-library test.",
        {
            "datetime": dt_str,
            "make": "Canon",
            "model": "EOS R5",
            "gps_lat": 51.5,
            "gps_lon": -0.1,
        },
    )

    # Reader 1: piexif
    exif1 = piexif.load(raw)
    dto_piexif = exif1["Exif"].get(piexif.ExifIFD.DateTimeOriginal, b"").decode()

    # Reader 2: exifread (independent library)
    tags = exifread.process_file(_io.BytesIO(raw), details=False)
    dto_exifread = str(tags.get("EXIF DateTimeOriginal", ""))

    assert dto_piexif == dt_str, f"piexif read {dto_piexif!r}, expected {dt_str!r}"
    assert dto_exifread == dt_str, f"exifread read {dto_exifread!r}, expected {dt_str!r}"
    assert dto_piexif == dto_exifread  # Both libraries agree


def test_jpeg_pillow_image_open(catalog: ProvenanceCatalog) -> None:
    """Pillow.Image.open() must be able to decode the produced JPEG bytes."""
    pytest.importorskip("PIL")
    import io as _io

    from PIL import Image

    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(
        ArtifactProfile.JPEG,
        "Pillow decode test.",
        {"datetime": "2024:02:20 08:00:00", "device": "cam02"},
    )
    img = Image.open(_io.BytesIO(raw))
    assert img.format == "JPEG"
    assert img.size == (320, 240)


# ---------------------------------------------------------------------------
# Log / CSV
# ---------------------------------------------------------------------------


def test_log_is_valid_csv(catalog: ProvenanceCatalog) -> None:
    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(
        ArtifactProfile.LOG,
        "User alice logged in.\nFile accessed: report.pdf.",
        {"source": "auth_service", "timestamp": "2024-01-15 08:00:00"},
    )
    assert isinstance(raw, bytes)
    decoded = raw.decode("utf-8")
    reader = csv.DictReader(io.StringIO(decoded))
    rows = list(reader)
    assert len(rows) >= 1
    assert "message" in (reader.fieldnames or [])
    assert "timestamp" in (reader.fieldnames or [])


def test_log_has_all_required_columns(catalog: ProvenanceCatalog) -> None:
    from evidence_factory.models import ArtifactProfile

    raw = catalog.write(
        ArtifactProfile.LOG,
        "Event logged.",
        {"source": "syslog", "timestamp": "2024-02-01 00:00:00"},
    )
    decoded = raw.decode("utf-8")
    reader = csv.DictReader(io.StringIO(decoded))
    required = {"timestamp", "level", "source", "message"}
    assert required.issubset(set(reader.fieldnames or []))
